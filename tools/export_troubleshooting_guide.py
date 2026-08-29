"""Export the Ops Hub troubleshooting guide to Excel.

Plain-language: what each diagnostic check means, what to do when it
fails, and whether it is safe to keep working. Written for the business
owner to use without waiting for help.

Usage:  python tools/export_troubleshooting_guide.py [output.xlsx]
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.app_paths import get_project_root

BRAND_GREEN = "1B7A3D"
FONT = "Arial"
ROOT = str(get_project_root())
APP_NAME = "Ops Hub"

CHECKS = [
    ("Project structure",
     f"One of the app's core folders or files is missing (assets, core, modules, database).",
     f"Almost always means files were moved or only part of the folder was copied. Restore the newest folder from {ROOT}\\backups, or re-copy the whole {APP_NAME} folder. Do not just recreate an empty folder - the contents are what matter.",
     "No - stop and restore first."),
    ("Python environment",
     "Informational only. Shows which Python is running the app.",
     "Nothing to fix. If you are running the .exe, this is the Python bundled inside it.",
     "Yes."),
    ("Required dependencies",
     "A library the app needs (customtkinter, pypdf, reportlab, openpyxl) is missing or too old.",
     "Only happens when running from source, never from the .exe. Run:  pip install -r requirements.txt",
     "Yes - use the .exe meanwhile."),
    ("Import health",
     "A core code file failed to load - usually a syntax error from a half-finished edit.",
     "The details line names the file. If this appeared right after a change, restore that file from backup or ask Claude. Your .exe is unaffected until it is rebuilt.",
     "Yes - use the .exe meanwhile."),
    ("Module registration",
     "One or more modules (CRM, Quotes, Accounting) did not register at startup.",
     "If a module folder was deleted or renamed, put it back. If it names one module, only that module is broken. NOTE: this reads FAIL if the check runs before the app has finished starting - reopen the Troubleshooter once the main window is up before believing it.",
     "Mostly - the named module will be missing."),
    ("Module discovery report",
     "The scan that finds module folders came back empty or errored.",
     f"Same causes and fix as Module registration. Check {ROOT}\\modules still has its subfolders and each contains a module.py.",
     "Only if it names specific modules."),
    ("Application folder access",
     "The app cannot write to its own folders (database, backups, exports).",
     f"Usually Windows permissions, or OneDrive locking the folder, or a file open in Excel. Close any open export and retry. If it persists, check {APP_NAME} is not inside a synced OneDrive folder.",
     "No - saving will fail."),
    ("Database health",
     "The database is corrupt, unreachable, or has integrity / foreign-key errors.",
     f"STOP. Do not keep entering data. Restore the newest {ROOT}\\database\\app.db.bak_* file over app.db, or restore a full backup folder. Tell Claude before doing anything else.",
     "No - stop immediately."),
    ("Backup services",
     "The backup system cannot reach its folder.",
     f"Make sure {ROOT}\\backups exists and is writable. Same OneDrive / permissions causes as folder access.",
     "Yes, but you are working unprotected - fix it soon."),
    ("Paths and resources",
     "The logo, fonts or other bundled assets cannot be found.",
     f"In a rebuilt .exe this usually means the assets folder was not bundled - documents would then print with no logo and in the wrong font. Ask Claude to rebuild. From source, check {ROOT}\\assets still exists.",
     "Yes, but check any PDF before sending it to a client."),
    ("Basic application health",
     "A roll-up of the checks above. Lists whatever is missing.",
     "Fix the individual checks it names; this one clears by itself once they do.",
     "Depends what it names."),
    ("Pending migrations (shown separately)",
     "The database is older than the app - new features need a schema update.",
     "Click Apply Pending Migrations in the Troubleshooter. It backs up first. If it errors, stop and tell Claude rather than retrying repeatedly.",
     "No - apply them first."),
]

SITUATIONS = [
    ("The app will not open at all",
     f"1. Try opening it once more - a stuck previous copy can block it.\n"
     f"2. Restart Windows.\n"
     f"3. Still failing: restore the newest folder from {ROOT}\\backups.\n"
     f"4. Tell Claude exactly what the error said."),
    ("The app opens but a module is blank or errors",
     "Open the Troubleshooter and run all checks - it will name the module. Everything else stays usable meanwhile."),
    ("Data looks wrong or missing",
     "Do NOT re-enter it yet. Check the Recycle Bin first - it may be deleted rather than lost. Then tell Claude. Restoring a backup over good data loses everything you have done since that backup."),
    ("A price is wrong on a quote",
     "Settings > Supplier Pricing. Edit the item's cost, then reopen the quote and update the unit price manually."),
    ("You changed a supplier price and nothing happened",
     "Prices apply to NEW quotes. Existing quotes keep their original figures - that is correct, not a bug."),
    ("A PDF prints with no logo or the wrong font",
     "The assets did not get bundled into the .exe. Ask Claude to rebuild it. Do not send that PDF to a client."),
    ("You want to undo something in Accounting",
     "The ledger has an Undo for the last edit. Use that rather than re-typing - it keeps the audit trail honest."),
    ("A bank statement imported without categories",
     "Categories come from saved rules. Change one transaction's category and say yes when asked to remember the rule - the next import will apply it automatically. See the Auto-Categorising sheet."),
    ("Before ANY big change",
     f"Make a backup: the Backup module, or copy the whole {APP_NAME} folder. Backups are free; recovery is not."),
]

AUTO_CATEGORY = [
    ("How it works",
     "When you change a transaction's category, the app asks two things: whether to fix every other transaction from the same shop, and whether to remember the rule for future imports. Saying yes to the second is what makes later statements categorise themselves."),
    ("Why a new statement can import uncategorised",
     "There is no saved rule for that shop yet. A brand-new payee the app has never seen cannot be categorised - it has nothing to go on. It is not broken."),
    ("How to teach it",
     "Import the statement, then set the category on one row per shop and say yes to remembering it. Each shop only needs teaching once. After a couple of statements most rows will come in already categorised."),
    ("How shops are matched",
     "The app strips card numbers and dates off the description, so KWIKSPAR HOMESTEAD 485442*7275 28 JUN and the 24 JUN one count as the same shop. Slight variations in the reference do not break it."),
    ("What it will never do",
     "It will not guess. A payee with no rule is left blank rather than being put in a category that might be wrong - a wrong category in your books is worse than an empty one."),
]


def _header(sheet, columns):
    for index, title in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BRAND_GREEN)
        cell.alignment = Alignment(horizontal="left", vertical="center")


def _put(sheet, row, column, value, bold=False):
    cell = sheet.cell(row=row, column=column, value=value)
    cell.font = Font(name=FONT, bold=bold)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    return cell


def _table(sheet, columns, rows, widths, heights):
    """Deliberately conservative formatting - column widths and wrapped
    text only. Explicit row heights and freeze panes were dropped after
    Excel refused to open an earlier version of this file; letting Excel
    size the rows itself is more portable and reads no worse."""

    _header(sheet, columns)
    for offset, record in enumerate(rows):
        row = offset + 2
        for index, value in enumerate(record, start=1):
            _put(sheet, row, index, value, bold=(index == 1))
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def main():
    destination = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("Ops_Hub_Troubleshooting_Guide.xlsx")

    workbook = openpyxl.Workbook()

    guide = workbook.active
    guide.title = "Troubleshooting Guide"
    _table(
        guide,
        ["Check", "What it means", "If it FAILS - what to do", "Can I keep working?"],
        CHECKS, [26, 44, 70, 26], 66,
    )

    _table(
        workbook.create_sheet("If All Else Fails"),
        ["Situation", "Do this"],
        SITUATIONS, [42, 96], 78,
    )

    _table(
        workbook.create_sheet("Auto-Categorising"),
        ["Question", "Answer"],
        AUTO_CATEGORY, [40, 98], 74,
    )

    log = workbook.create_sheet("Health Check Log")
    _table(
        log,
        ["Date", "Run by", "Result", "Notes / what was fixed"],
        [("2026-08-29", "Claude", "All 11 checks PASS",
          "Updated baseline for Ops Hub. DB created fresh, all checks PASS. Modules: CRM, Quotes, Accounting, Accounting Workbook, Calendar, Communications, Documents, Gallery, Projects, Proposals, Backup, Settings, Dashboard. Site Visit removed (not applicable).")],
        [14, 16, 24, 84], 34,
    )

    workbook.save(destination)
    print(f"Wrote {destination}")


if __name__ == "__main__":
    main()
