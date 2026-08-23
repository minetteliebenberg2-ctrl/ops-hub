"""Job Card docx export.

Same header/footer design as Quote/Invoice PDFs (Lato, brand green #21A94D,
FacilitiesCo logo, business block). No amounts, no remittance — work-log only.
"""

import shutil
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from core.app_paths import get_assets_dir

GREEN = RGBColor(0x21, 0xA9, 0x4D)
GREEN_HEX = "21A94D"
CHARCOAL = RGBColor(0x33, 0x33, 0x33)
GREY = RGBColor(0x70, 0x70, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Lato"

ASSETS_DIR = get_assets_dir()
LOGO_PATH = ASSETS_DIR / "logo_placeholder.png"

BUSINESS_LINE1 = "Shade Solutions by FacilitiesCo (Pty) Ltd"
BUSINESS_LINE2 = "Germiston, South Africa, 1401"
BUSINESS_LINE3 = "sales@facilitiesco.com"
BUSINESS_LINE4 = "010 015 0532 / 083 378 5122"
BUSINESS_LINE5 = "www.facilitiesco.com"


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------

def _set_font(run, size=10, bold=False, color=CHARCOAL, italic=False):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = FONT
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), FONT)


def _shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _cell_padding(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)


def _row_height(row, cm):
    trPr = row._tr.get_or_add_trPr()
    trH = OxmlElement("w:trHeight")
    trH.set(qn("w:val"), str(int(cm * 567)))
    trH.set(qn("w:hRule"), "atLeast")
    trPr.append(trH)


def _no_border(table):
    """Remove all table borders."""
    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        tblBorders.append(el)
    existing = tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        tblPr.remove(existing)
    tblPr.append(tblBorders)


def _paragraph(parent, text="", size=10, bold=False, color=CHARCOAL,
               italic=False, align=None, space_before=0, space_after=4):
    p = parent.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if align:
        p.alignment = align
    if text:
        r = p.add_run(text)
        _set_font(r, size=size, bold=bold, color=color, italic=italic)
    return p


def _green_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), GREEN_HEX)
    pBdr.append(bottom)
    pPr.append(pBdr)


# ---------------------------------------------------------------------------
# Header
# ---------------------------------------------------------------------------

def _build_header(doc):
    """Logo left, business block right, separated by a green rule."""
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _no_border(table)
    row = table.rows[0]
    _row_height(row, 2.8)

    # Logo cell
    logo_cell = row.cells[0]
    logo_cell.width = Cm(8)
    logo_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _cell_padding(logo_cell, 0, 0, 0, 0)
    p = logo_cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if LOGO_PATH.exists():
        run = p.add_run()
        run.add_picture(str(LOGO_PATH), width=Cm(7))
    else:
        r = p.add_run("FacilitiesCo")
        _set_font(r, size=16, bold=True, color=GREEN)

    # Business block cell
    biz_cell = row.cells[1]
    biz_cell.width = Cm(9.5)
    biz_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _cell_padding(biz_cell, 0, 0, 0, 0)

    lines = [
        (BUSINESS_LINE1, 9.5, True),
        (BUSINESS_LINE2, 9, False),
        (BUSINESS_LINE3, 9, False),
        (BUSINESS_LINE4, 9, False),
        (BUSINESS_LINE5, 9, False),
    ]
    first = True
    for text, size, bold in lines:
        if first:
            p = biz_cell.paragraphs[0]
            first = False
        else:
            p = biz_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(text)
        _set_font(r, size=size, bold=bold, color=CHARCOAL)

    _green_rule(doc)


# ---------------------------------------------------------------------------
# Work-log table
# ---------------------------------------------------------------------------

COLUMNS = ["Date", "Type", "Job", "Description", "Staff", "Completed", "Director"]
# widths in cm, total must fit within A4 body (~17 cm)
COL_WIDTHS = [2.2, 2.2, 2.2, 4.8, 2.0, 2.0, 1.6]


def _build_work_log_table(doc, entries):
    table = doc.add_table(rows=1, cols=len(COLUMNS))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _no_border(table)

    # Header row
    hdr = table.rows[0]
    _row_height(hdr, 0.8)
    for i, col in enumerate(COLUMNS):
        cell = hdr.cells[i]
        cell.width = Cm(COL_WIDTHS[i])
        _shade_cell(cell, GREEN_HEX)
        _cell_padding(cell, 40, 40, 80, 80)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        r = p.add_run(col)
        _set_font(r, size=9, bold=True, color=WHITE)

    # Data rows
    for entry in entries:
        row = table.add_row()
        _row_height(row, 0.75)
        completed_txt = "✓" if getattr(entry, "completed", False) else ""
        vals = [
            getattr(entry, "entry_date", "") or "",
            getattr(entry, "work_type", "") or "",
            getattr(entry, "job_summary", "") or "",
            getattr(entry, "description", "") or "",
            getattr(entry, "staff", "") or "",
            completed_txt,
            getattr(entry, "director_signoff", "") or "",
        ]
        for i, val in enumerate(vals):
            cell = row.cells[i]
            cell.width = Cm(COL_WIDTHS[i])
            _cell_padding(cell, 40, 40, 80, 80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            # light bottom rule between rows
            tcPr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "4")
            bottom.set(qn("w:space"), "0")
            bottom.set(qn("w:color"), "E0E0E0")
            borders.append(bottom)
            tcPr.append(borders)
            p = cell.paragraphs[0]
            r = p.add_run(str(val))
            _set_font(r, size=9, color=CHARCOAL)


# ---------------------------------------------------------------------------
# Footer
# ---------------------------------------------------------------------------

def _build_footer(doc):
    _paragraph(doc, "", space_before=8, space_after=0)
    _green_rule(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("DESIGN  |  CREATE  |  INNOVATE  |  MAINTAIN")
    _set_font(r, size=8, color=GREY, italic=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(1)
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run("Reg No. 2024/772013/07  ·  @facilitiesco")
    _set_font(r2, size=8, color=GREY)


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------

def generate_job_card_docx(job_card, customer, entries, output_path):
    """Write a Job Card docx to output_path.

    Backs up the previous file as <output_path>.prev.docx before overwriting.
    Returns the final path.
    """
    output_path = Path(output_path)

    if output_path.exists():
        prev = output_path.with_suffix(".prev.docx")
        shutil.copy2(output_path, prev)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    # Default style
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), FONT)

    # ---- Header ----
    _build_header(doc)

    # ---- Title ----
    _paragraph(doc, "JOB CARD", size=18, bold=True,
               align=WD_ALIGN_PARAGRAPH.CENTER,
               space_before=4, space_after=10)

    # ---- Meta grid ----
    issue_date = (job_card.created_at or "")[:10] or str(date.today())
    meta_left = [
        ("Job Card #", job_card.job_card_number or "—"),
        ("Date", issue_date),
        ("Customer ID", customer.customer_number or customer.id[:8]),
    ]
    bill_to = job_card.bill_to_name or customer.name

    meta_table = doc.add_table(rows=len(meta_left), cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta_table.autofit = False
    _no_border(meta_table)

    for i, (label, value) in enumerate(meta_left):
        left_cell = meta_table.rows[i].cells[0]
        left_cell.width = Cm(8.75)
        _cell_padding(left_cell, 20, 20, 0, 0)
        lp = left_cell.paragraphs[0]
        lp.paragraph_format.space_after = Pt(2)
        r1 = lp.add_run(f"{label}:  ")
        _set_font(r1, size=9, bold=True)
        r2 = lp.add_run(value)
        _set_font(r2, size=9)

        right_cell = meta_table.rows[i].cells[1]
        right_cell.width = Cm(8.75)
        _cell_padding(right_cell, 20, 20, 40, 0)
        rp = right_cell.paragraphs[0]
        rp.paragraph_format.space_after = Pt(2)
        if i == 0:
            r3 = rp.add_run("Bill To:  ")
            _set_font(r3, size=9, bold=True)
            r4 = rp.add_run(bill_to)
            _set_font(r4, size=9)

    # ---- Delivery address ----
    delivery = (job_card.delivery_address or "").strip()
    if delivery:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run("Delivery Address:  ")
        _set_font(r, size=9, bold=True)
        r2 = p.add_run(delivery)
        _set_font(r2, size=9)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ---- Work-log table ----
    _build_work_log_table(doc, entries)

    # ---- Footer ----
    _build_footer(doc)

    doc.save(str(output_path))
    return output_path
