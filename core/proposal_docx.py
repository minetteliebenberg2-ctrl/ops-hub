# ==========================================================
# FC Hub - Proposal DOCX Generator
# ----------------------------------------------------------
# Purpose:
# Fill a generic proposal layout with real client/site data from a
# ProposalData object. Every "[ FILL IN ]" style field becomes data-driven, falling back
# to the same placeholder text when a field is left blank.
# ==========================================================

import io
import os
from datetime import datetime

from docx import Document
from PIL import Image as _PILImage, ExifTags as _ExifTags
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from core.app_paths import get_assets_dir

GREEN = RGBColor(0x1B, 0x7A, 0x3D)
CHARCOAL = RGBColor(0x33, 0x33, 0x33)
GREY = RGBColor(0x70, 0x70, 0x70)
LIGHTGREY = RGBColor(0xAA, 0xAA, 0xAA)
PLACEHOLDER_BG = "FAFAFA"
FONT = "Poppins"

ASSETS_DIR = get_assets_dir()
LOGO_PATH = ASSETS_DIR / "logo_placeholder.png"
ICON_DIR = ASSETS_DIR / "social_icons"

TIMELINE_PHASES = [
    ("1. Initial Assessment", "Requirements confirmed with the client"),
    ("2. Quotation", "Formal quote issued for sign-off"),
    ("3. Order Confirmation", "Signed quotation + deposit received"),
    ("4. Preparation", "Materials and resources prepared"),
    ("5. Delivery", "Work carried out as quoted"),
    ("6. Final Sign-off", "Sign-off with the client contact"),
]


def _business_settings():
    try:
        from core.business_settings_service import BusinessSettingsService
        return BusinessSettingsService().get_settings()
    except Exception:
        return None


def _contact_lines(settings):
    """Business name + contact lines from Settings (no hardcoded identity)."""
    if settings is None:
        return "", []
    name = (settings.trading_name or settings.legal_name or "").strip()
    contact = "  ·  ".join(v.strip() for v in (settings.email, settings.phone) if (v or "").strip())
    lines = [contact] if contact else []
    if (settings.website or "").strip():
        lines.append(settings.website.strip())
    return name, lines


def _val(value, placeholder):
    """Return the real value if set, else the template's placeholder text."""
    value = (value or "").strip()
    return value if value else placeholder


_MAX_PX = 1000
_JPEG_QUALITY = 85


def _auto_orient(img):
    """Rotate a PIL Image to match its EXIF orientation tag, then strip it."""
    try:
        exif = img._getexif()
        if exif:
            for tag, value in exif.items():
                if _ExifTags.TAGS.get(tag) == "Orientation":
                    if value == 3:
                        img = img.rotate(180, expand=True)
                    elif value == 6:
                        img = img.rotate(270, expand=True)
                    elif value == 8:
                        img = img.rotate(90, expand=True)
                    break
    except (AttributeError, KeyError, IndexError):
        pass
    return img


def _compress_image(path, manual_rotation=0):
    """Return a BytesIO JPEG at max 1000x1000 px, quality 85.
    Applies EXIF auto-orient first, then any manual rotation (degrees CW).
    Keeps originals untouched; falls back to the raw path on any error."""
    try:
        with _PILImage.open(path) as img:
            img = _auto_orient(img)
            if manual_rotation:
                img = img.rotate(-manual_rotation, expand=True)
            img = img.convert("RGB")
            img.thumbnail((_MAX_PX, _MAX_PX), _PILImage.Resampling.LANCZOS)
            buf = io.BytesIO()
            img.save(buf, "JPEG", quality=_JPEG_QUALITY)
            buf.seek(0)
            return buf
    except Exception:
        return path


def set_font(run, size=10.5, color=CHARCOAL, bold=False, italic=False, name=FONT):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def cell_border(cell, style="dashed", color="AAAAAA", size=6):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), style)
        el.set(qn('w:sz'), str(size))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    tcPr.append(borders)


def set_row_height(row, cm_height):
    trPr = row._tr.get_or_add_trPr()
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), str(int(cm_height * 567)))
    trHeight.set(qn('w:hRule'), 'atLeast')
    trPr.append(trHeight)


def set_bottom_border(paragraph, color="1B7A3D", size=8, space=5):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), str(space))
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def page_break(doc):
    doc.add_page_break()


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text.upper())
    set_font(run, size=19, color=CHARCOAL, bold=True)
    set_bottom_border(p)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text.upper())
    set_font(run, size=12, color=GREEN, bold=True)
    return p


def body(doc, text, size=10.5, color=CHARCOAL, italic=False, bold=False, space_after=6, alignment=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if alignment:
        p.alignment = alignment
    run = p.add_run(text)
    set_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_font(run, size=size, color=CHARCOAL)
    return p


def field_line(doc, label, value, placeholder="[ FILL IN ]"):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r1 = p.add_run(f"{label}:  ")
    set_font(r1, size=10.5, color=CHARCOAL, bold=True)
    text = _val(value, placeholder)
    is_placeholder = text == placeholder
    r2 = p.add_run(text)
    set_font(r2, size=10.5, color=LIGHTGREY if is_placeholder else CHARCOAL)
    return p


def cell_padding(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for edge, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:w'), str(value))
        el.set(qn('w:type'), 'dxa')
        tcMar.append(el)
    tcPr.append(tcMar)


def row_bottom_rule(cell, color="E0E0E0", size=4):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), '0')
    bottom.set(qn('w:color'), color)
    borders.append(bottom)
    tcPr.append(borders)


def styled_table(doc, headers, rows, col_widths_cm):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr_cells = table.rows[0].cells
    set_row_height(table.rows[0], 0.9)
    for i, htext in enumerate(headers):
        hdr_cells[i].width = Cm(col_widths_cm[i])
        shade_cell(hdr_cells[i], "1B7A3D")
        cell_padding(hdr_cells[i])
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = hdr_cells[i].paragraphs[0]
        r = p.add_run(htext)
        set_font(r, size=10, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
    for row_index, row_data in enumerate(rows):
        row_cells = table.add_row().cells
        set_row_height(table.rows[row_index + 1], 0.85)
        is_last = row_index == len(rows) - 1
        for i, val in enumerate(row_data):
            row_cells[i].width = Cm(col_widths_cm[i])
            cell_padding(row_cells[i])
            row_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if not is_last:
                row_bottom_rule(row_cells[i])
            p = row_cells[i].paragraphs[0]
            r = p.add_run(str(val))
            set_font(r, size=10, color=CHARCOAL)
    return table


def photo_grid(doc, photo_paths, count=8, cols=4, cell_w_cm=3.7, cell_h_cm=2.3, rotations=None):
    rows_needed = (count + cols - 1) // cols
    table = doc.add_table(rows=rows_needed, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    n = 0
    for r in range(rows_needed):
        set_row_height(table.rows[r], cell_h_cm)
        for c in range(cols):
            n += 1
            cell = table.rows[r].cells[c]
            cell.width = Cm(cell_w_cm)
            photo_path = photo_paths[n - 1] if n - 1 < len(photo_paths) else None
            if photo_path and os.path.exists(photo_path):
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                try:
                    angle = (rotations or {}).get(photo_path, 0)
                    run.add_picture(_compress_image(photo_path, manual_rotation=angle), width=Cm(cell_w_cm - 0.3))
                except Exception:
                    shade_cell(cell, PLACEHOLDER_BG)
                    cell_border(cell)
            else:
                shade_cell(cell, PLACEHOLDER_BG)
                cell_border(cell)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r1 = p.add_run("\U0001F4F7")
                set_font(r1, size=14, color=RGBColor(0xBB, 0xBB, 0xBB))
                p2 = cell.add_paragraph()
                p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r2 = p2.add_run(f"Photo {n}" if n <= count else "")
                set_font(r2, size=8, color=RGBColor(0xBB, 0xBB, 0xBB), italic=True)
    return table


def add_logo(doc, path, width_cm, alignment=WD_ALIGN_PARAGRAPH.CENTER):
    p = doc.add_paragraph()
    p.alignment = alignment
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm))
    return p


def generate_proposal_docx(data, output_path, include_terms=False, photo_rotations=None):
    """Build a filled proposal document from a ProposalData
    object and save it to output_path."""

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal_style = doc.styles['Normal']
    normal_style.font.name = FONT
    normal_style.font.size = Pt(10.5)
    rpr = normal_style.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), FONT)

    settings = _business_settings()
    business_name, contact_lines = _contact_lines(settings)

    # ---------------- PAGE 1 - COVER ----------------
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    if LOGO_PATH.exists():
        add_logo(doc, str(LOGO_PATH), width_cm=9)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    p.paragraph_format.space_after = Pt(40)
    r = p.add_run("PROPOSAL")
    set_font(r, size=38, color=CHARCOAL, bold=True)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.width = Cm(7)
    set_row_height(table.rows[0], 3.2)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    logo_path = data.client_logo_path
    if logo_path and os.path.exists(logo_path):
        cell_border(cell, style="single", color="E5E5E5", size=4)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(_compress_image(logo_path), width=Cm(5.5))
    else:
        shade_cell(cell, PLACEHOLDER_BG)
        cell_border(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("[ CLIENT LOGO ]")
        set_font(r, size=11, color=LIGHTGREY, italic=True)

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    field_line(doc, "Prepared For", data.client.company_name, "[CLIENT / COMPANY NAME]")
    field_line(doc, "Attention", data.attention, "[CONTACT PERSON]")
    field_line(doc, "Site", data.site, "[SITE NAME / ADDRESS]")
    field_line(doc, "Date", data.proposal_date.split("T")[0] if data.proposal_date else "", datetime.now().strftime("%Y-%m-%d"))

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(90)

    if business_name:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"Prepared by {business_name}")
        set_font(r, size=10.5, color=GREY)
    for line in contact_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        set_font(r, size=9.5, color=GREY)

    page_break(doc)

    # ---------------- PAGE 2 - SCOPE OF WORK ----------------
    h1(doc, "Scope of Work")

    h2(doc, "Site Details")
    field_line(doc, "Site / Area", data.site_area, "[SITE / AREA]")
    field_line(doc, "Dimensions / Size", data.dimensions, "[DIMENSIONS]")

    h2(doc, "Project Timeline")
    durations = list(data.timeline_durations or [])
    durations += [""] * (6 - len(durations))
    styled_table(
        doc,
        ["Phase", "Description", "Duration"],
        [
            [phase, desc, _val(durations[i], "[ ]")]
            for i, (phase, desc) in enumerate(TIMELINE_PHASES)
        ],
        [4.8, 8.2, 3.0],
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    h2(doc, "What Is Expected From You")
    bullet(doc, "A deposit is required on confirmation of order, with the balance payable on completion, as set out in the quotation")
    bullet(doc, "Confirmation of order is a signed quotation and the deposit")
    bullet(doc, "Access arranged for assessment, delivery, and final sign-off")

    # ---------------- SITE PHOTOS ----------------
    if data.site_photos:
        page_break(doc)
        h1(doc, "Site Photos")
        body(doc, f"Area: {_val(data.site_area, '[SITE / AREA NAME]')}", bold=True, size=12, color=CHARCOAL, space_after=8)
        photo_grid(doc, data.site_photos[:8], count=8, cols=4, rotations=photo_rotations)
        if len(data.site_photos) > 8:
            page_break(doc)
            h1(doc, "Additional Site Photos")
            photo_grid(doc, data.site_photos[8:16], count=8, cols=4, rotations=photo_rotations)

    page_break(doc)

    # ---------------- CLOSING / CONTACT ----------------
    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Thank You for the Opportunity")
    set_font(r, size=24, color=CHARCOAL, bold=True)

    for _ in range(3):
        doc.add_paragraph()

    if business_name:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(business_name)
        set_font(r, size=12.5, color=CHARCOAL, bold=True)
    for line in contact_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        set_font(r, size=10, color=GREY)

    doc.save(output_path)
    return output_path
