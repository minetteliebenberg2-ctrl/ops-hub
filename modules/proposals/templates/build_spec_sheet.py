import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "..")))

import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from core.app_paths import get_assets_dir

GREEN = RGBColor(0x1B, 0x7A, 0x3D)
CHARCOAL = RGBColor(0x33, 0x33, 0x33)
GREY = RGBColor(0x70, 0x70, 0x70)
FONT = "Poppins"

DOCS_DIR = os.path.dirname(os.path.abspath(__file__))
ASSETS_DIR = get_assets_dir()
LOGO_PATH = ASSETS_DIR / "logo_placeholder.png"
ICON_DIR = ASSETS_DIR / "social_icons"


def set_font(run, size=10.5, color=CHARCOAL, bold=False, italic=False, name=FONT):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def set_bottom_border(paragraph, color="1B7A3D", size=8, space=5):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), str(space))
    bottom.set(qn('w:color'), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text.upper())
    set_font(run, size=12, color=GREEN, bold=True)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_font(run, size=10.5, color=CHARCOAL)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.6)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

normal_style = doc.styles['Normal']
normal_style.font.name = FONT
normal_style.font.size = Pt(10.5)
rpr = normal_style.element.get_or_add_rPr()
rFonts = rpr.find(qn('w:rFonts'))
if rFonts is None:
    rFonts = OxmlElement('w:rFonts')
    rpr.append(rFonts)
rFonts.set(qn('w:eastAsia'), FONT)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
run = p.add_run()
run.add_picture(str(LOGO_PATH), width=Cm(7))

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
run = p.add_run("TECHNICAL SPECIFICATION SHEET")
set_font(run, size=18, color=CHARCOAL, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
run = p.add_run("SHADEPORT SYSTEMS")
set_font(run, size=12, color=GREEN, bold=True)
set_bottom_border(p, size=10)

section_heading(doc, "Structural Steel Framework — Cantilevers")
bullet(doc, "Anchor Poles: 152mm / 165mm round tubing (3mm wall thickness — heavy-duty)")
bullet(doc, "Main Frame: 50mm / 57mm round tubing (2mm wall thickness)")

section_heading(doc, "Structural Steel Framework — Four Post")
bullet(doc, "Anchor Poles: 76mm / 101mm round tubing (3mm wall thickness — heavy-duty)")
bullet(doc, "Main Frame: 50mm round tubing (2mm wall thickness)")

section_heading(doc, "Foundations & Anchoring")
bullet(doc, "600mm – 800mm deep pole planting")
bullet(doc, "Secured with ready-mix concrete")

section_heading(doc, "Wind & Load Performance")
bullet(doc, "Balanced cantilever load distribution")
bullet(doc, "Designed for high wind conditions (120–140 km/h when correctly installed)")

section_heading(doc, "Shade Netting")
bullet(doc, "Knittex Z25 UV-stabilised shade netting")
bullet(doc, "High durability and tear resistance")

section_heading(doc, "Cable Tensioning System (Certified)")
bullet(doc, "5mm 6x19 Fibre Core Galvanised Steel Cable")
bullet(doc, "Breaking Load: 13.6 kN (≈ 1,385 kg)")
bullet(doc, "Tested Breaking Load: 14.3 kN")
bullet(doc, "Secured with Crosby clamps")

section_heading(doc, "Stitching & Thread — Coats Dabond AWF")
bullet(doc, "Bonded, twisted continuous filament polyester thread")
bullet(doc, "PFC-free anti-wick finish blocks water migrating through stitched seams — a prerequisite for genuinely water-repellent seams, not just weatherproof fabric")
bullet(doc, "Dual-level lubrication protects against needle heat and keeps high-speed sewing smooth")
bullet(doc, "“Z”-twist construction, purpose-built for single-needle high-speed industrial sewing")
bullet(doc, "Excellent bleach, mildew and rot resistance for long-term outdoor exposure")
bullet(doc, "Two tex weights specified by seam load: Tex 80 (5,700 cN breaking strength) for standard seams, Tex 135 (9,310 cN breaking strength) for heavy-duty seams")

section_heading(doc, "Installation Standard")
bullet(doc, "Precision alignment of cantilever arms")
bullet(doc, "Correct structural positioning")
bullet(doc, "Professional tensioning and finishing")

section_heading(doc, "System Summary")
bullet(doc, "Heavy-duty steel framework")
bullet(doc, "Deep concrete foundations")
bullet(doc, "High-strength certified cable system")
bullet(doc, "Coats-certified stitching for water- and weather-resistant seams")
bullet(doc, "Built for long-term durability and performance")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(20)
pPr = p._p.get_or_add_pPr()
pBdr = OxmlElement('w:pBdr')
top = OxmlElement('w:top')
top.set(qn('w:val'), 'single')
top.set(qn('w:sz'), '4')
top.set(qn('w:space'), '8')
top.set(qn('w:color'), 'CCCCCC')
pBdr.append(top)
pPr.append(pBdr)


def footer_line(text, bold=False, size=10.5, color=GREY):
    fp = doc.add_paragraph()
    fp.paragraph_format.space_after = Pt(1)
    r = fp.add_run(text)
    set_font(r, size=size, color=color, bold=bold)
    return fp


footer_line("Prepared By: Minette Liebenberg", bold=True, size=11, color=CHARCOAL)
footer_line("Director")
footer_line("Shade Solutions by FacilitiesCo", bold=True)
footer_line("sales@facilitiesco.com  ·  010 015 0532 / 083 378 5122")

doc.add_paragraph().paragraph_format.space_after = Pt(10)

p = doc.add_paragraph()
for icon_name in ["instagram", "facebook", "youtube", "linkedin"]:
    run = p.add_run()
    run.add_picture(str(ICON_DIR / f"icon_{icon_name}.png"), width=Cm(0.9))
    p.add_run("   ")

doc.add_paragraph().paragraph_format.space_after = Pt(4)
footer_line("www.facilitiesco.com  ·  @facilitiesco", size=9.5, color=GREY)

out_path = os.path.join(DOCS_DIR, "FacilitiesCo_Technical_Spec_Sheet.docx")
doc.save(out_path)
print("saved:", out_path)
