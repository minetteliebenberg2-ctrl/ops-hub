"""Tests for the Documents module.

The behaviour that matters most here is expiry: a Tax Compliance PIN or
COIDA letter that lapsed unnoticed is the failure this module exists to
prevent, so the date handling gets the most attention.
"""

import sqlite3
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

from modules.documents.services import (
    CATEGORIES,
    EXPIRY_WARNING_DAYS,
    Document,
    DocumentsRepository,
    parse_date,
)


def _build_schema(path):
    """Minimal schema matching migration v0017."""

    connection = sqlite3.connect(str(path))
    connection.executescript(
        """
        CREATE TABLE business_settings (
            id TEXT PRIMARY KEY,
            updated_at TEXT NOT NULL DEFAULT '',
            documents_root TEXT NOT NULL DEFAULT ''
        );
        INSERT INTO business_settings (id) VALUES ('business');
        CREATE TABLE documents (
            id TEXT PRIMARY KEY,
            title TEXT NOT NULL,
            category TEXT NOT NULL DEFAULT 'Compliance',
            stored_filename TEXT NOT NULL,
            original_filename TEXT NOT NULL DEFAULT '',
            issue_date TEXT NOT NULL DEFAULT '',
            expiry_date TEXT NOT NULL DEFAULT '',
            notes TEXT NOT NULL DEFAULT '',
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            updated_by TEXT NOT NULL DEFAULT ''
        );
        """
    )
    connection.commit()
    connection.close()


class DateParsingTests(unittest.TestCase):

    def test_accepts_the_formats_users_actually_type(self):
        for value in ("2027-03-31", "2027/03/31", "31/03/2027", "31-03-2027"):
            self.assertEqual(parse_date(value), date(2027, 3, 31), value)

    def test_blank_and_nonsense_return_none_rather_than_raising(self):
        for value in ("", "   ", None, "not a date", "2027-13-45"):
            self.assertIsNone(parse_date(value))


class ExpiryTests(unittest.TestCase):

    def _document(self, expiry):
        return Document(title="Tax Compliance", expiry_date=expiry)

    def test_document_without_expiry_never_reports_expired(self):
        document = self._document("")
        self.assertIsNone(document.days_until_expiry())
        self.assertFalse(document.is_expired)
        self.assertFalse(document.is_expiring_soon)

    def test_lapsed_document_is_expired_not_merely_expiring(self):
        yesterday = (date.today() - timedelta(days=1)).isoformat()
        document = self._document(yesterday)
        self.assertTrue(document.is_expired)
        self.assertFalse(document.is_expiring_soon)
        self.assertEqual(document.days_until_expiry(), -1)

    def test_document_inside_warning_window_is_expiring_soon(self):
        soon = (date.today() + timedelta(days=EXPIRY_WARNING_DAYS - 1)).isoformat()
        document = self._document(soon)
        self.assertTrue(document.is_expiring_soon)
        self.assertFalse(document.is_expired)

    def test_document_beyond_warning_window_is_neither(self):
        later = (date.today() + timedelta(days=EXPIRY_WARNING_DAYS + 10)).isoformat()
        document = self._document(later)
        self.assertFalse(document.is_expiring_soon)
        self.assertFalse(document.is_expired)

    def test_expiry_boundary_day_still_counts_as_expiring_not_expired(self):
        """A certificate expiring today is still valid today."""
        document = self._document(date.today().isoformat())
        self.assertEqual(document.days_until_expiry(), 0)
        self.assertTrue(document.is_expiring_soon)
        self.assertFalse(document.is_expired)


class RepositoryTests(unittest.TestCase):

    def setUp(self):
        self._temp = tempfile.TemporaryDirectory()
        self.temp_path = Path(self._temp.name)
        self.database_path = self.temp_path / "test.db"
        _build_schema(self.database_path)

        self.root = self.temp_path / "documents"
        self.repository = DocumentsRepository(database_path=self.database_path)
        self.repository.set_documents_root(self.root)

        self.source = self.temp_path / "COIDA.pdf"
        self.source.write_bytes(b"%PDF-1.4 fake")

    def tearDown(self):
        self._temp.cleanup()

    def test_adding_copies_the_file_and_leaves_the_original_alone(self):
        document = self.repository.add_document(
            source_path=self.source, title="COIDA 2026", category="Compliance"
        )
        self.assertTrue(self.source.is_file(), "the user's original must not be moved")
        stored = self.repository.file_path(document)
        self.assertTrue(stored.is_file())
        self.assertEqual(stored.read_bytes(), b"%PDF-1.4 fake")

    def test_same_filename_added_twice_does_not_overwrite(self):
        first = self.repository.add_document(source_path=self.source, title="First")
        second = self.repository.add_document(source_path=self.source, title="Second")
        self.assertNotEqual(first.stored_filename, second.stored_filename)
        self.assertTrue(self.repository.file_path(first).is_file())
        self.assertTrue(self.repository.file_path(second).is_file())

    def test_unknown_category_is_rejected(self):
        with self.assertRaises(ValueError):
            self.repository.add_document(
                source_path=self.source, title="X", category="Nonsense"
            )

    def test_missing_source_file_is_rejected(self):
        with self.assertRaises(FileNotFoundError):
            self.repository.add_document(
                source_path=self.temp_path / "nope.pdf", title="X"
            )

    def test_categories_cover_what_the_module_offers(self):
        self.assertEqual(set(CATEGORIES), {"Compliance", "Letters", "Banking", "Insurance"})

    def test_listing_can_filter_by_category(self):
        self.repository.add_document(source_path=self.source, title="A", category="Compliance")
        self.repository.add_document(source_path=self.source, title="B", category="Insurance")
        self.assertEqual(len(self.repository.list_documents()), 2)
        insurance = self.repository.list_documents("Insurance")
        self.assertEqual([d.title for d in insurance], ["B"])

    def test_expiring_documents_are_ordered_most_urgent_first(self):
        today = date.today()
        self.repository.add_document(
            source_path=self.source, title="Lapsed",
            expiry_date=(today - timedelta(days=10)).isoformat())
        self.repository.add_document(
            source_path=self.source, title="Soon",
            expiry_date=(today + timedelta(days=5)).isoformat())
        self.repository.add_document(
            source_path=self.source, title="Fine",
            expiry_date=(today + timedelta(days=400)).isoformat())
        self.repository.add_document(source_path=self.source, title="Never expires")

        titles = [d.title for d in self.repository.expiring_documents()]
        self.assertEqual(titles, ["Lapsed", "Soon"])

    def test_deleting_removes_the_stored_copy_but_not_the_original(self):
        document = self.repository.add_document(source_path=self.source, title="Temp")
        stored = self.repository.file_path(document)
        self.assertTrue(self.repository.delete_document(document.id))
        self.assertFalse(stored.is_file())
        self.assertTrue(self.source.is_file(), "the user's original must survive")
        self.assertIsNone(self.repository.get(document.id))

    def test_update_only_touches_allowed_fields(self):
        document = self.repository.add_document(source_path=self.source, title="Old")
        updated = self.repository.update_document(
            document.id, title="New", expiry_date="2027-01-01",
            stored_filename="hacked.pdf",
        )
        self.assertEqual(updated.title, "New")
        self.assertEqual(updated.expiry_date, "2027-01-01")
        self.assertEqual(updated.stored_filename, document.stored_filename)

    def test_documents_root_is_created_and_configurable(self):
        # get_documents_root() is what creates the folder, so that a freshly
        # configured location works without the user making it first.
        self.assertEqual(self.repository.get_documents_root(), self.root)
        self.assertTrue(self.root.is_dir())


class TemplateTests(unittest.TestCase):
    """The templates must not leak the home address or the retired tagline."""

    def setUp(self):
        self.template_dir = Path(__file__).parents[1] / "modules" / "documents" / "templates"

    def test_letterhead_exists_and_is_branded(self):
        path = self.template_dir / "FacilitiesCo_Letterhead.docx"
        self.assertTrue(path.is_file(), "run build_templates.py to generate it")

        from docx import Document as DocxDocument

        document = DocxDocument(str(path))
        section = document.sections[0]
        header_text = " ".join(
            cell.text for table in section.header.tables
            for row in table.rows for cell in row.cells
        )
        footer_text = " ".join(p.text for p in section.footer.paragraphs)

        self.assertIn("Shade Solutions by FacilitiesCo", header_text)
        self.assertIn("Germiston, South Africa, 1401", header_text)
        self.assertIn("MAINTAIN", footer_text)

        everything = header_text + footer_text + " ".join(p.text for p in document.paragraphs)
        self.assertNotIn("SUSTAIN", everything)
        self.assertNotIn("Francis", everything, "home address must never appear")

    def test_spreadsheet_exists_and_is_branded(self):
        path = self.template_dir / "FacilitiesCo_Spreadsheet.xlsx"
        self.assertTrue(path.is_file(), "run build_templates.py to generate it")

        from openpyxl import load_workbook

        sheet = load_workbook(str(path)).active
        values = " ".join(
            str(sheet.cell(row=r, column=c).value or "")
            for r in range(1, 12) for c in range(1, 8)
        )
        self.assertIn("Shade Solutions by FacilitiesCo", values)
        self.assertIn("Germiston, South Africa, 1401", values)
        self.assertIn("MAINTAIN", values)
        self.assertNotIn("SUSTAIN", values)
        self.assertNotIn("Francis", values)


if __name__ == "__main__":
    unittest.main()
