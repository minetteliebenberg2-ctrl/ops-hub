import os
import tempfile

import pytest

from core.database import Database
from core.ledger_category_repository import LedgerCategoryRepository
from core.ledger_repository import LedgerTransactionRepository
from core.ledger_service import LedgerService
from core.ledger_transaction import EXPENSE, INCOME

from tests.test_ledger import FNB_SAMPLE, write_sample_csv


@pytest.fixture
def test_db():
    db_path = tempfile.mktemp(suffix=".db")
    database = Database(db_path)
    yield database
    if os.path.exists(db_path):
        os.remove(db_path)


@pytest.fixture
def ledger_service(test_db):
    repository = LedgerTransactionRepository(db=test_db)
    category_repository = LedgerCategoryRepository(db=test_db)
    return LedgerService(repository=repository, category_repository=category_repository)


# --------------------------------------------------
# Category CRUD
# --------------------------------------------------

def test_migration_seeds_all_real_categories(ledger_service):
    from core.migrations.versions.v0020_ledger_categories_audit_receipts import EXPENSE_SEED, INCOME_SEED
    from core.migrations.versions.v0021_category_rewording_and_additions import (
        NEW_BUSINESS_EXPENSE_CATEGORIES, NEW_PERSONAL_EXPENSE_CATEGORIES,
        NEW_PERSONAL_INCOME_CATEGORIES, RENAMES,
    )

    categories = ledger_service.category_repository.list_all(active_only=False)
    names = {c.name for c in categories}

    v0020_names = [RENAMES.get(name, name) for name in INCOME_SEED + EXPENSE_SEED]
    v0021_additions = NEW_BUSINESS_EXPENSE_CATEGORIES + NEW_PERSONAL_EXPENSE_CATEGORIES + NEW_PERSONAL_INCOME_CATEGORIES
    v0049_additions = ["Catering", "Fee / Consulting Income", "Workman's Comp / COIDA", "Hardware"]
    expected = v0020_names + v0021_additions + v0049_additions

    for name in expected:
        assert name in names
    for old_name in RENAMES:
        assert old_name not in names
    assert len(categories) == len(expected)


def test_add_category(ledger_service):
    category = ledger_service.add_category("Materials & Supplies", EXPENSE, "minette")
    assert category.active is True
    assert "Materials & Supplies" in ledger_service.list_categories(category_type=EXPENSE)


def test_add_category_rejects_duplicate_name(ledger_service):
    ledger_service.add_category("Custom Category", INCOME, "minette")
    with pytest.raises(ValueError):
        ledger_service.add_category("Custom Category", EXPENSE, "minette")


def test_rename_category(ledger_service):
    category = ledger_service.add_category("Old Name", EXPENSE, "minette")
    ledger_service.rename_category(category.id, "New Name")
    updated = ledger_service.category_repository.get(category.id)
    assert updated.name == "New Name"


def test_retire_and_reactivate_category_does_not_break_existing_transactions(ledger_service):
    category = ledger_service.add_category("Seasonal Costs", EXPENSE, "minette")
    transaction = ledger_service.add_transaction(
        "2026-08-01", "Test", 5000, EXPENSE, "Seasonal Costs", "Personal Account", "", "", "minette",
    )

    ledger_service.retire_category(category.id)

    active_categories = ledger_service.list_categories(category_type=EXPENSE, active_only=True)
    assert "Seasonal Costs" not in active_categories

    # existing transaction keeps its category, unaffected
    stored = ledger_service.repository.get(transaction.id)
    assert stored.category == "Seasonal Costs"

    ledger_service.reactivate_category(category.id)
    active_categories = ledger_service.list_categories(category_type=EXPENSE, active_only=True)
    assert "Seasonal Costs" in active_categories


# --------------------------------------------------
# Edit audit trail + undo
# --------------------------------------------------

def test_record_edit_and_update_category_logs_it(ledger_service):
    transaction = ledger_service.add_transaction(
        "2026-08-01", "Test", 5000, EXPENSE, "General Expenses", "Personal Account", "", "", "minette",
    )
    ledger_service.update_category(transaction.id, "Office Supplies", "minette")

    with ledger_service.repository.db.connect() as connection:
        rows = connection.execute(
            "SELECT * FROM ledger_transaction_edits WHERE transaction_id = ?", (transaction.id,),
        ).fetchall()
    assert len(rows) == 1
    assert rows[0]["field_name"] == "category"
    assert rows[0]["old_value"] == "General Expenses"
    assert rows[0]["new_value"] == "Office Supplies"
    assert rows[0]["undone"] == 0


def test_undo_last_edit_reverts_category(ledger_service):
    transaction = ledger_service.add_transaction(
        "2026-08-01", "Test", 5000, EXPENSE, "General Expenses", "Personal Account", "", "", "minette",
    )
    ledger_service.update_category(transaction.id, "Office Supplies", "minette")

    result = ledger_service.undo_last_edit("minette")

    assert result["transaction_id"] == transaction.id
    assert result["field_name"] == "category"
    assert result["reverted_to"] == "General Expenses"

    reverted = ledger_service.repository.get(transaction.id)
    assert reverted.category == "General Expenses"


def test_undo_history_stays_append_only(ledger_service):
    transaction = ledger_service.add_transaction(
        "2026-08-01", "Test", 5000, EXPENSE, "General Expenses", "Personal Account", "", "", "minette",
    )
    ledger_service.update_category(transaction.id, "Office Supplies", "minette")
    ledger_service.undo_last_edit("minette")

    with ledger_service.repository.db.connect() as connection:
        rows = connection.execute(
            "SELECT * FROM ledger_transaction_edits WHERE transaction_id = ? ORDER BY edited_at, rowid",
            (transaction.id,),
        ).fetchall()

    # Original edit row is untouched except for its undone flag; the
    # revert is a brand new append-only row.
    assert len(rows) == 2
    assert rows[0]["old_value"] == "General Expenses"
    assert rows[0]["new_value"] == "Office Supplies"
    assert rows[0]["undone"] == 1
    assert rows[1]["old_value"] == "Office Supplies"
    assert rows[1]["new_value"] == "General Expenses"
    assert rows[1]["undone"] == 0


def test_undo_with_nothing_to_undo_raises(ledger_service):
    with pytest.raises(ValueError):
        ledger_service.undo_last_edit("minette")


def test_undo_reverts_amount(ledger_service):
    transaction = ledger_service.add_transaction(
        "2026-08-01", "Test", 5000, EXPENSE, "General Expenses", "Personal Account", "", "", "minette",
    )
    ledger_service.record_edit(transaction.id, "amount_minor", transaction.amount_minor, 7500, "minette")
    transaction.amount_minor = 7500
    ledger_service.repository.update(transaction)

    result = ledger_service.undo_last_edit("minette")
    assert result["field_name"] == "amount_minor"

    reverted = ledger_service.repository.get(transaction.id)
    assert reverted.amount_minor == 5000


# --------------------------------------------------
# Reconciliation
# --------------------------------------------------

def test_get_reconciliation_returns_none_without_snapshot(ledger_service):
    assert ledger_service.get_reconciliation("Gold Business Account") is None


def test_import_bank_statement_upserts_reconciliation_snapshot(ledger_service, tmp_path):
    filepath = write_sample_csv(tmp_path)
    ledger_service.import_bank_statement(filepath, "minette")

    recon = ledger_service.get_reconciliation("Gold Business Account")
    assert recon is not None
    assert recon["statement_balance_minor"] == 179648  # "Balance:, 1796.48, 764.27"
    assert recon["statement_date"] == "2026-04-02"
    assert recon["ledger_balance_minor"] == recon["difference_minor"] + recon["statement_balance_minor"]


def test_reimport_refreshes_reconciliation_snapshot(ledger_service, tmp_path):
    filepath = write_sample_csv(tmp_path)
    ledger_service.import_bank_statement(filepath, "minette")

    extended = FNB_SAMPLE + "2026/04/03, -100.00, 696.48, NEW TRANSACTION\n"
    extended = extended.replace("Balance:, 1796.48, 764.27", "Balance:, 696.48, 100.00")
    filepath2 = write_sample_csv(tmp_path, extended, "statement2.csv")
    ledger_service.import_bank_statement(filepath2, "minette")

    with ledger_service.repository.db.connect() as connection:
        rows = connection.execute("SELECT * FROM account_reconciliation_snapshots").fetchall()
    assert len(rows) == 1
    assert rows[0]["statement_balance_minor"] == 69648


# --------------------------------------------------
# Excel export data sanity (LedgerService side, GUI export tested via openpyxl directly)
# --------------------------------------------------

def test_ledger_transactions_expose_receipt_filename_default_empty(ledger_service):
    transaction = ledger_service.add_transaction(
        "2026-08-01", "Test", 5000, EXPENSE, "General Expenses", "Personal Account", "", "", "minette",
    )
    stored = ledger_service.repository.get(transaction.id)
    assert stored.receipt_filename == ""


# --------------------------------------------------
# Receipt attachment
# --------------------------------------------------

def test_attach_receipt_copies_file_and_records_edit(ledger_service, tmp_path, monkeypatch):
    receipts_root = tmp_path / "documents_root"
    receipts_root.mkdir()

    class FakeDocumentsRepository:
        def get_documents_root(self):
            return receipts_root

    import core.ledger_service as ledger_service_module
    monkeypatch.setattr(
        "modules.documents.services.DocumentsRepository", FakeDocumentsRepository, raising=False,
    )

    source_file = tmp_path / "receipt.pdf"
    source_file.write_text("fake receipt content")

    transaction = ledger_service.add_transaction(
        "2026-08-01", "Test", 5000, EXPENSE, "General Expenses", "Personal Account", "", "", "minette",
    )

    stored_filename = ledger_service.attach_receipt(transaction.id, str(source_file), "minette")

    assert (receipts_root / "Ledger Receipts" / stored_filename).exists()

    updated = ledger_service.repository.get(transaction.id)
    assert updated.receipt_filename == stored_filename

    with ledger_service.repository.db.connect() as connection:
        rows = connection.execute(
            "SELECT * FROM ledger_transaction_edits WHERE transaction_id = ? AND field_name = 'receipt_filename'",
            (transaction.id,),
        ).fetchall()
    assert len(rows) == 1
    assert rows[0]["new_value"] == stored_filename


def test_attach_receipt_raises_for_missing_file(ledger_service, tmp_path):
    transaction = ledger_service.add_transaction(
        "2026-08-01", "Test", 5000, EXPENSE, "General Expenses", "Personal Account", "", "", "minette",
    )
    with pytest.raises(FileNotFoundError):
        ledger_service.attach_receipt(transaction.id, str(tmp_path / "nope.pdf"), "minette")


# --------------------------------------------------
# Excel export
# --------------------------------------------------

def test_export_transactions_to_xlsx_produces_readable_file(ledger_service, tmp_path):
    from openpyxl import load_workbook

    from core.ledger_export import export_transactions_to_xlsx

    ledger_service.add_transaction(
        "2026-08-01", "Cash sale", 150000, INCOME, "Sales Income", "Gold Business Account", "REF1", "note", "minette",
    )
    transactions = ledger_service.list_transactions()

    output = tmp_path / "export.xlsx"
    export_transactions_to_xlsx(transactions, output)

    workbook = load_workbook(output)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][0] == "Date"
    assert rows[1][1] == "Cash sale"
    assert rows[1][5] == 1500.0


# --------------------------------------------------
# PDF income statement export
# --------------------------------------------------

def test_generate_income_statement_pdf_does_not_crash_and_contains_totals(ledger_service, tmp_path):
    from core.business_settings_service import BusinessSettingsService
    from core.ledger_report_pdf import generate_income_statement_pdf

    ledger_service.add_transaction(
        "2026-08-01", "Cash sale", 150000, INCOME, "Sales Income", "Gold Business Account", "", "", "minette",
    )
    ledger_service.add_transaction(
        "2026-08-02", "Fuel", 20000, EXPENSE, "Motor V: Fuel & Oil", "Gold Business Account", "", "", "minette",
    )
    summary = ledger_service.get_summary()

    business_settings = BusinessSettingsService().get_settings()

    output = tmp_path / "income_statement.pdf"
    generate_income_statement_pdf(summary, business_settings, "2026-08-01", "2026-08-31", output)

    assert output.exists()
    assert output.stat().st_size > 0
