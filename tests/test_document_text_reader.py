"""Regression tests for reading text out of Word and PDF documents."""

import pytest

from core.document_text_reader import DocumentTextReader, UnsupportedDocumentType


@pytest.fixture
def sample_docx(tmp_path):
    import docx

    path = tmp_path / "sample.docx"
    document = docx.Document()
    document.add_paragraph("Bright Facilities (Pty) Ltd")
    document.add_paragraph("011 234 5678")
    document.save(str(path))
    return path


@pytest.fixture
def sample_pdf(tmp_path):
    from reportlab.pdfgen import canvas

    path = tmp_path / "sample.pdf"
    c = canvas.Canvas(str(path))
    c.drawString(100, 750, "Acme Cleaning CC")
    c.save()
    return path


def test_reads_text_from_docx(sample_docx):
    text = DocumentTextReader().read(sample_docx)
    assert "Bright Facilities (Pty) Ltd" in text
    assert "011 234 5678" in text


def test_reads_text_from_pdf(sample_pdf):
    text = DocumentTextReader().read(sample_pdf)
    assert "Acme Cleaning CC" in text


def test_missing_file_raises_clear_error(tmp_path):
    missing = tmp_path / "does_not_exist.docx"
    with pytest.raises(FileNotFoundError):
        DocumentTextReader().read(missing)


def test_unsupported_extension_raises_clear_error(tmp_path):
    bad_file = tmp_path / "notes.txt"
    bad_file.write_text("just some text")
    with pytest.raises(UnsupportedDocumentType):
        DocumentTextReader().read(bad_file)
